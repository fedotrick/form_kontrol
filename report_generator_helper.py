import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
import pandas as pd
from datetime import datetime
import os

class ReportHelper:
    def __init__(self, template_path="Отчёт.xlsx"):
        """
        Инициализирует помощник по работе с отчетами.
        
        Args:
            template_path (str): Путь к шаблону отчета
        """
        self.template_path = template_path
        self.workbook = None
        self.worksheet = None
        
        # Индексы строк с шапками
        self.header1_row = 7  # Первая шапка (Доработка) 
        self.header2_row = 18  # Вторая шапка (Окончательный брак и Второй сорт)
        
        # Индексы строк для заполнения данных
        self.data1_start_row = 8  # Начало данных после первой шапки
        self.data2_start_row = 19  # Начало данных после второй шапки
        
        # Загрузим шаблон
        self._load_template()
        
        # Словари для маппинга колонок
        self.process_headers()
    
    def _load_template(self):
        """Загружает шаблон отчета"""
        if not os.path.exists(self.template_path):
            raise FileNotFoundError(f"Шаблон отчета не найден: {self.template_path}")
        
        self.workbook = openpyxl.load_workbook(self.template_path)
        self.worksheet = self.workbook.active
    
    def process_headers(self):
        """Обрабатывает заголовки таблиц и создает маппинги колонок"""
        # Словари для хранения маппингов колонок
        self.header1_mapping = {}  # Маппинг для первой таблицы (доработка)
        self.header2_mapping = {}  # Маппинг для второй таблицы (брак и второй сорт)
        
        # Обработка первой шапки (доработка)
        for cell in self.worksheet[self.header1_row]:
            if cell.value:
                self.header1_mapping[cell.value] = cell.column_letter
        
        # Обработка второй шапки (брак и второй сорт)
        for cell in self.worksheet[self.header2_row]:
            if cell.value:
                self.header2_mapping[cell.value] = cell.column_letter
        
        print(f"Маппинг первой шапки (строка {self.header1_row}):")
        for field, col in self.header1_mapping.items():
            print(f"  {field}: колонка {col}")
        
        print(f"Маппинг второй шапки (строка {self.header2_row}):")
        for field, col in self.header2_mapping.items():
            print(f"  {field}: колонка {col}")
    
    def find_main_cell_in_merge(self, row, column):
        """
        Находит основную ячейку в объединенном диапазоне.
        
        Args:
            row (int): Номер строки
            column (int): Номер столбца
            
        Returns:
            tuple: Координаты основной ячейки (row, column) или None, если ячейка не объединена
        """
        cell_coord = f"{openpyxl.utils.get_column_letter(column)}{row}"
        for merged_range in self.worksheet.merged_cells.ranges:
            min_col, min_row, max_col, max_row = openpyxl.utils.range_boundaries(merged_range.coord)
            if (min_row <= row <= max_row) and (min_col <= column <= max_col):
                return (min_row, min_col)
        return None
    
    def set_date(self, date_str=None):
        """
        Устанавливает дату в отчете
        
        Args:
            date_str (str, optional): Строка с датой. Если None, используется текущая дата.
        """
        if date_str is None:
            date_str = datetime.now().strftime("%d.%m.%Y")
        
        print(f"Устанавливаем дату в отчете: {date_str}")
        
        # Извлекаем год из date_str
        try:
            date_obj = datetime.strptime(date_str, "%d.%m.%Y")
            year = date_obj.year
        except:
            # Если не получилось извлечь год, используем текущий
            year = datetime.now().year
            print(f"Не удалось распознать дату {date_str}, используем текущий год: {year}")
        
        # Находим ячейку с датой (строка 1-3, начинается с "Дата:")
        date_cell_found = False
        for row_idx in range(1, 4):
            for col_idx in range(1, 5):
                cell = self.worksheet.cell(row=row_idx, column=col_idx)
                if cell.value and isinstance(cell.value, str) and "Дата:" in cell.value:
                    # Находим дату в самой строке и модифицируем ее
                    current_value = cell.value
                    updated_value = f"Дата:  {date_str}"
                    
                    # Обновляем значение в ячейке
                    main_coords = self.find_main_cell_in_merge(row_idx, col_idx)
                    if main_coords:
                        self.worksheet.cell(row=main_coords[0], column=main_coords[1]).value = updated_value
                    else:
                        cell.value = updated_value
                    
                    date_cell_found = True
                    print(f"Обновлена ячейка даты: {updated_value}")
                    return True
        
        # Если не найдена ячейка с датой, создаем новую
        if not date_cell_found:
            self.worksheet.cell(row=1, column=2).value = f"Дата: {date_str}"
            print(f"Создана новая ячейка с датой: {date_str}")
        
        return True
    
    def set_controllers(self, controllers):
        """
        Устанавливает имена контролеров в отчете
        
        Args:
            controllers (list): Список имен контролеров
        """
        controllers_str = ", ".join(controllers)
        
        # Находим ячейку с контролерами (строка 1-3, содержит "Контролеры:")
        for row_idx in range(1, 4):
            for col_idx in range(1, 20):
                cell = self.worksheet.cell(row=row_idx, column=col_idx)
                if cell.value and isinstance(cell.value, str) and "Контролеры:" in cell.value:
                    # Вместо попытки изменить соседнюю ячейку, изменяем саму ячейку
                    current_value = cell.value
                    updated_value = f"Контролеры: {controllers_str}"
                    
                    # Обновляем значение в ячейке
                    main_coords = self.find_main_cell_in_merge(row_idx, col_idx)
                    if main_coords:
                        self.worksheet.cell(row=main_coords[0], column=main_coords[1]).value = updated_value
                    else:
                        cell.value = updated_value
                    
                    return True
        
        # Если не найдена ячейка с контролерами, создаем новую
        self.worksheet.cell(row=1, column=15).value = f"Контролеры: {controllers_str}"
        return True
    
    def add_data_row_1(self, row_data, row_index=None):
        """
        Добавляет строку данных в первую таблицу (доработка)
        
        Args:
            row_data (dict): Словарь с данными для добавления
            row_index (int, optional): Индекс строки. Если None, добавляется новая строка.
        
        Returns:
            int: Индекс добавленной строки
        """
        if row_index is None:
            # Находим первую пустую строку
            row_index = self.data1_start_row
            while self.worksheet.cell(row=row_index, column=1).value:
                row_index += 1
                if row_index >= self.header2_row:  # Не пересекаем границу со второй таблицей
                    break
        
        # Заполняем данные в соответствии с маппингом
        for field, value in row_data.items():
            if field in self.header1_mapping:
                col = self.header1_mapping[field]
                col_idx = openpyxl.utils.column_index_from_string(col)
                # Определяем основную ячейку, если она объединена
                main_coords = self.find_main_cell_in_merge(row_index, col_idx)
                if main_coords:
                    self.worksheet.cell(row=main_coords[0], column=main_coords[1]).value = value
                else:
                    self.worksheet.cell(row=row_index, column=col_idx).value = value
        
        return row_index
    
    def add_data_row_2(self, row_data, row_index=None):
        """
        Добавляет строку данных во вторую таблицу (брак и второй сорт)
        
        Args:
            row_data (dict): Словарь с данными для добавления
            row_index (int, optional): Индекс строки. Если None, добавляется новая строка.
        
        Returns:
            int: Индекс добавленной строки
        """
        if row_index is None:
            # Находим первую пустую строку
            row_index = self.data2_start_row
            while self.worksheet.cell(row=row_index, column=1).value:
                row_index += 1
                if row_index >= self.worksheet.max_row:  # Не выходим за пределы листа
                    break
        
        # Заполняем данные в соответствии с маппингом
        for field, value in row_data.items():
            if field in self.header2_mapping:
                col = self.header2_mapping[field]
                col_idx = openpyxl.utils.column_index_from_string(col)
                # Определяем основную ячейку, если она объединена
                main_coords = self.find_main_cell_in_merge(row_index, col_idx)
                if main_coords:
                    self.worksheet.cell(row=main_coords[0], column=main_coords[1]).value = value
                else:
                    self.worksheet.cell(row=row_index, column=col_idx).value = value
        
        return row_index
    
    def save(self, output_path=None):
        """
        Сохраняет отчет в указанный файл
        
        Args:
            output_path (str, optional): Путь для сохранения. Если None, используется шаблон имени с датой.
        
        Returns:
            str: Путь к сохраненному файлу
        """
        if output_path is None:
            # Пытаемся найти дату в отчете, чтобы использовать ее в имени файла
            date_found = None
            for row_idx in range(1, 4):
                for col_idx in range(1, 5):
                    cell = self.worksheet.cell(row=row_idx, column=col_idx)
                    if cell.value and isinstance(cell.value, str) and "Дата:" in cell.value:
                        # Извлекаем дату из ячейки
                        try:
                            date_part = cell.value.split("Дата:")[1].strip()
                            # Преобразуем формат даты из DD.MM.YYYY в DD-MM-YYYY для имени файла
                            date_obj = datetime.strptime(date_part, "%d.%m.%Y")
                            date_found = date_obj.strftime("%d-%m-%Y")
                            print(f"Извлечена дата для имени файла: {date_found}")
                        except:
                            print(f"Не удалось извлечь дату из ячейки: {cell.value}")
                            pass
            
            # Если дата не найдена, используем текущую
            if not date_found:
                date_found = datetime.now().strftime("%d-%m-%Y")
                print(f"Используем текущую дату для имени файла: {date_found}")
            
            output_path = f"Отчет_{date_found}.xlsx"
        
        # Проверяем, существует ли директория
        dir_path = os.path.dirname(output_path)
        if dir_path and not os.path.exists(dir_path):
            os.makedirs(dir_path)
            print(f"Создана директория: {dir_path}")
        
        # Получаем абсолютный путь к файлу
        abs_path = os.path.abspath(output_path)
        
        # Сохраняем файл
        self.workbook.save(output_path)
        print(f"Отчет сохранен в файл: {abs_path}")
        
        return output_path
    
    def load_from_control_xlsx(self, control_file="control.xlsx", date_str=None):
        """
        Загружает данные из файла control.xlsx и добавляет их в отчет
        
        Args:
            control_file (str): Путь к файлу с данными контроля
            date_str (str, optional): Дата для фильтрации (DD.MM.YYYY)
        
        Returns:
            bool: True если данные успешно загружены
        """
        if not os.path.exists(control_file):
            print(f"Файл {control_file} не найден")
            return False
            
        try:
            print(f"Загрузка данных из {control_file} для даты {date_str}")
            df = pd.read_excel(control_file)
            
            # Заменяем NaN значения на пустые строки для строковых полей
            string_columns = df.select_dtypes(include=['object']).columns
            df[string_columns] = df[string_columns].fillna('')
            
            # Заменяем NaN значения на 0 для числовых полей
            numeric_columns = df.select_dtypes(include=['float64', 'int64']).columns
            df[numeric_columns] = df[numeric_columns].fillna(0)
            
            # Фильтруем по дате, если указана
            if date_str:
                # Прямое сравнение строк, так как даты хранятся как строки
                if 'Контроль_дата_приемки' in df.columns:
                    print(f"Фильтрация строк с датой {date_str}")
                    # Сначала проверяем точное соответствие строк
                    df_filtered = df[df['Контроль_дата_приемки'] == date_str]
                    
                    # Если не нашли точное соответствие, пробуем преобразовать
                    if len(df_filtered) == 0:
                        print(f"Точное соответствие дат не найдено, пробуем преобразование...")
                        # Преобразуем столбец с датой в datetime для универсального сравнения
                        df['Контроль_дата_приемки'] = pd.to_datetime(df['Контроль_дата_приемки'], errors='coerce')
                        date_obj = datetime.strptime(date_str, "%d.%m.%Y").date()
                        df_filtered = df[df['Контроль_дата_приемки'].dt.date == date_obj]
                    
                    print(f"Найдено {len(df_filtered)} записей с датой {date_str}")
                    if len(df_filtered) == 0:
                        print(f"Записи с датой {date_str} не найдены в файле {control_file}")
                        return False
                        
                    df = df_filtered
                    
            # Убеждаемся, что все необходимые колонки есть в dataframe
            required_columns = ['Наименование_отливки']
            missing_columns = [col for col in required_columns if col not in df.columns]
            if missing_columns:
                print(f"Отсутствуют обязательные колонки: {', '.join(missing_columns)}")
                return False
                
            # Проверяем на пустые наименования отливок
            empty_names = df[df['Наименование_отливки'] == ''].shape[0]
            if empty_names > 0:
                print(f"Внимание: {empty_names} записей с пустыми наименованиями отливок")
                
                # Пытаемся восстановить наименования из файла plavka.xlsx, если он существует
                if os.path.exists('plavka.xlsx'):
                    print("Пытаемся восстановить наименования из файла plavka.xlsx...")
                    try:
                        df_plavka = pd.read_excel('plavka.xlsx')
                        # Заменяем None на пустые строки
                        df_plavka = df_plavka.fillna('')
                        
                        # Для каждой записи с пустым наименованием пытаемся найти соответствующую запись в plavka.xlsx
                        for idx, row in df[df['Наименование_отливки'] == ''].iterrows():
                            plavka_number = row['Номер_плавки']
                            matching_rows = df_plavka[df_plavka['Учетный_номер'].astype(str) == str(plavka_number)]
                            
                            if not matching_rows.empty and 'Наименование_отливки' in matching_rows.columns:
                                name = matching_rows['Наименование_отливки'].iloc[0]
                                if name and name != '':
                                    df.at[idx, 'Наименование_отливки'] = name
                                    print(f"Восстановлено наименование для номера плавки {plavka_number}: {name}")
                        
                        # Пересчитываем количество пустых наименований
                        empty_names = df[df['Наименование_отливки'] == ''].shape[0]
                        if empty_names > 0:
                            print(f"После восстановления остались {empty_names} записей с пустыми наименованиями")
                    except Exception as e:
                        print(f"Ошибка при попытке восстановить наименования: {str(e)}")

            # Группировка по наименованиям отливок
            print("Группировка данных по наименованиям отливок...")
            
            # Создаем соответствие исходных колонок для группировки
            grouped_by_name = df.groupby('Наименование_отливки')
            
            # Для каждой группы (наименования отливки) создаем строки в отчете
            for name, group in grouped_by_name:
                if name == '':  # Пропускаем пустые наименования
                    continue
                
                # Данные для первой таблицы (доработка)
                row_data1 = {"Наименование_отливки": name}
                
                # Суммируем данные по отлито
                row_data1["Контроль_отлито"] = int(group["Контроль_отлито"].sum()) if "Контроль_отлито" in group.columns else 0
                
                # Заполняем данные по доработке
                for col in df.columns:
                    if col.startswith("Доработка_") and col in self.header1_mapping:
                        row_data1[col] = int(group[col].sum()) if col in group.columns else 0
                
                # Добавляем строку в первую таблицу
                self.add_data_row_1(row_data1)
                
                # Данные для второй таблицы (брак и второй сорт)
                row_data2 = {"Наименование_отливки": name}
                
                # Суммируем данные по принято
                row_data2["Контроль_принято"] = int(group["Контроль_принято"].sum()) if "Контроль_принято" in group.columns else 0
                
                # Заполняем данные по окончательному браку
                for col in df.columns:
                    if col.startswith("Окончательный_брак_") and col in self.header2_mapping:
                        row_data2[col] = int(group[col].sum()) if col in group.columns else 0
                
                # Заполняем данные по второму сорту
                for col in df.columns:
                    if col.startswith("Второй_сорт_") and col in self.header2_mapping:
                        row_data2[col] = int(group[col].sum()) if col in group.columns else 0
                
                # Формируем список номеров плавок
                if "Номер_плавки" in group.columns:
                    unique_plavki = group["Номер_плавки"].unique()
                    row_data2["Номер_плавки"] = ", ".join(str(p) for p in unique_plavki if pd.notna(p) and str(p).strip())
                
                # Добавляем строку во вторую таблицу
                self.add_data_row_2(row_data2)
                
            print(f"Данные успешно загружены и сгруппированы по {len(grouped_by_name)} наименованиям отливок")
            return True
            
        except Exception as e:
            print(f"Ошибка при загрузке данных из {control_file}: {str(e)}")
            import traceback
            traceback.print_exc()
            return False
    
    def get_controllers_from_excel(self, control_file="control.xlsx", date_str=None):
        """
        Извлекает список контролеров из файла control.xlsx
        
        Args:
            control_file (str): Путь к файлу control.xlsx
            date_str (str, optional): Строка с датой для фильтрации. Если None, берутся все записи.
            
        Returns:
            list: Список имен контролеров
        """
        if not os.path.exists(control_file):
            print(f"Файл {control_file} не найден")
            return []
            
        try:
            print(f"Извлечение контролеров из {control_file} для даты {date_str}")
            df = pd.read_excel(control_file)
            
            # Фильтруем по дате, если указана
            if date_str:
                # Прямое сравнение строк, так как даты хранятся как строки
                if 'Контроль_дата_приемки' in df.columns:
                    # Сначала проверяем точное соответствие строк
                    df_filtered = df[df['Контроль_дата_приемки'] == date_str]
                    
                    # Если не нашли точное соответствие, пробуем преобразовать
                    if len(df_filtered) == 0:
                        print(f"Точное соответствие дат не найдено, пробуем преобразование...")
                        # Преобразуем столбец с датой в datetime для универсального сравнения
                        df['Контроль_дата_приемки'] = pd.to_datetime(df['Контроль_дата_приемки'], errors='coerce')
                        date_obj = datetime.strptime(date_str, "%d.%m.%Y").date()
                        df_filtered = df[df['Контроль_дата_приемки'].dt.date == date_obj]
                    
                    print(f"Найдено {len(df_filtered)} записей с контролерами за дату {date_str}")
                    if len(df_filtered) == 0:
                        return []
                        
                    df = df_filtered
                    
            # Собираем имена контролеров
            controllers = set()
            controller_columns = ['Контролер1', 'Контролер2', 'Контролер3']
            
            for column in controller_columns:
                if column in df.columns:
                    # Заменяем NaN на пустые строки
                    df[column] = df[column].fillna('')
                    # Собираем уникальные непустые имена
                    controllers.update(name for name in df[column].unique() if name and isinstance(name, str))
                    
            print(f"Извлечено {len(controllers)} уникальных контролеров")
            return sorted(list(controllers))
            
        except Exception as e:
            print(f"Ошибка при извлечении контролеров: {str(e)}")
            import traceback
            traceback.print_exc()
            return []

# Пример использования
if __name__ == "__main__":
    try:
        # Создаем экземпляр помощника
        helper = ReportHelper()
        
        # Устанавливаем дату
        current_date = datetime.now().strftime("%d.%m.%Y")
        helper.set_date(current_date)
        
        # Проверяем наличие файла control.xlsx
        if os.path.exists("control.xlsx"):
            # Извлекаем контролеров из файла
            controllers = helper.get_controllers_from_excel("control.xlsx", current_date)
            if controllers:
                print(f"Автоматически извлечены контролеры: {', '.join(controllers)}")
                helper.set_controllers(controllers)
            else:
                print("Контролеры не найдены в файле, используем значения по умолчанию")
                helper.set_controllers(["Елхова", "Лабуткина"])
            
            # Загружаем данные
            helper.load_from_control_xlsx("control.xlsx", current_date)
        else:
            print("Файл control.xlsx не найден, используем тестовые данные")
            
            # Устанавливаем контролеров вручную
            helper.set_controllers(["Елхова", "Лабуткина"])
            
            # Данные для первой таблицы (доработка)
            data1 = {
                "Наименование_отливки": "Вороток",
                "Контроль_отлито": 100,
                "Доработка_раковины": 5,
                "Доработка_зарез": 3,
                "Доработка_несоответствие_размеров": 2
            }
            helper.add_data_row_1(data1)
            
            # Еще одна строка для первой таблицы
            data1_2 = {
                "Наименование_отливки": "Ригель",
                "Контроль_отлито": 150,
                "Доработка_раковины": 7,
                "Доработка_зарез": 4,
                "Доработка_вырыв": 2
            }
            helper.add_data_row_1(data1_2)
            
            # Данные для второй таблицы (брак и второй сорт)
            data2 = {
                "Наименование_отливки": "Вороток",
                "Контроль_принято": 90,
                "Окончательный_брак_раковины": 3,
                "Окончательный_брак_коробление": 2,
                "Номер_плавки": "5-107/25"
            }
            helper.add_data_row_2(data2)
            
            # Еще одна строка для второй таблицы
            data2_2 = {
                "Наименование_отливки": "Ригель",
                "Контроль_принято": 135,
                "Окончательный_брак_вырыв": 4,
                "Окончательный_брак_спай": 1,
                "Номер_плавки": "5-132/25"
            }
            helper.add_data_row_2(data2_2)
        
        # Сохраняем отчет
        helper.save("Пример_отчета.xlsx")
        print("Готово!")
    
    except Exception as e:
        print(f"Произошла ошибка: {e}")
        import traceback
        traceback.print_exc() 